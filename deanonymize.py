# -*- coding: utf-8 -*-
"""
deanonymize.py - przywrocenie oryginalnych danych po anonimizacji.

Uzycie:
    python deanonymize.py                              (autodetekcja)
    python deanonymize.py input_anon.xlsx output.xlsx
"""

import sys
import json
import re
import time
from pathlib import Path
import openpyxl

MAP_FILE = "anon_map.json"
TOKEN_RE = re.compile(r"\[(OSOBA|FIRMA|MARKA|MIEJSCE|DANE|KWOTA)_\d+\]")


def deanonymize(input_path: str, output_path: str):
    t0 = time.time()
    print(f"Wczytuje mape: {MAP_FILE}")

    if not Path(MAP_FILE).exists():
        print(f"BLAD: Nie znaleziono {MAP_FILE}")
        sys.exit(1)

    with open(MAP_FILE, encoding="utf-8") as f:
        deanon_map = json.load(f)

    token_to_original = deanon_map.get("token_to_original", {})
    original_cells    = deanon_map.get("original_cells", [])
    print(f"  Tokenow: {len(token_to_original)}, komorek w mapie: {len(original_cells)}")

    cell_lookup = {
        (ci["sheet"], ci["row"], ci["col"]): (ci["original"], ci["anonymized"])
        for ci in original_cells
    }

    # token->oryginal, sortowane od najdluzszych
    replace_map = {
        token: originals[0]
        for token, originals in sorted(
            token_to_original.items(), key=lambda x: len(x[0]), reverse=True
        )
        if originals
    }

    print(f"Wczytuje: {input_path}")
    t1 = time.time()
    is_xlsm = input_path.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(input_path, keep_vba=is_xlsm)
    print(f"  Zaladowano w {time.time()-t1:.1f}s")

    restored = 0
    token_replaced = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if not cell.value or not isinstance(cell.value, str):
                    continue
                val = cell.value
                key = (sheet_name, cell.row, cell.column)

                if key in cell_lookup:
                    orig_val, anon_val = cell_lookup[key]
                    if val == anon_val:
                        cell.value = orig_val
                        restored += 1
                        continue

                if TOKEN_RE.search(val):
                    new_val = val
                    for token, orig in replace_map.items():
                        if token in new_val:
                            new_val = new_val.replace(token, orig)
                    if new_val != val:
                        cell.value = new_val
                        token_replaced += 1

    print(f"Zapisuje: {output_path}")
    wb.save(output_path)

    # Sprawdz pozostale tokeny
    remaining = []
    wb2 = openpyxl.load_workbook(output_path, read_only=True)
    for sn in wb2.sheetnames:
        for row in wb2[sn].iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and TOKEN_RE.search(cell.value):
                    col_letter = openpyxl.utils.get_column_letter(cell.column)
                    remaining.append(f"    {sn}!{col_letter}{cell.row}: {TOKEN_RE.findall(cell.value)}")
    wb2.close()

    print(f"\nGotowe! ({time.time()-t0:.1f}s)")
    print(f"  Przywrocono bezposrednio: {restored}")
    print(f"  Podmieniono tokeny:       {token_replaced}")
    print(f"  Plik wynikowy: {output_path}")

    if remaining:
        print(f"\n  [UWAGA] Nieodtworzone tokeny ({len(remaining)} komorek):")
        for ex in remaining[:10]:
            print(ex)
        if len(remaining) > 10:
            print(f"    ... i {len(remaining)-10} wiecej")
        print("  To prawdopodobnie nowe dane dodane przez Claude.")


if __name__ == "__main__":
    if len(sys.argv) >= 3:
        deanonymize(sys.argv[1], sys.argv[2])
    else:
        found = [p for p in Path(".").iterdir()
                 if p.stem.endswith("_anon")
                 and p.suffix.lower() in (".xls", ".xlsm", ".xlsx")
                 and p.is_file()
                 and not p.name.startswith("~$")]
        if not found:
            found = [p for p in Path(".").iterdir()
                     if p.suffix.lower() in (".xls", ".xlsm", ".xlsx")
                     and p.is_file()
                     and "entities" not in p.name
                     and not p.name.startswith("~$")]
        if not found:
            print("Nie znaleziono pliku .xlsx.")
            sys.exit(1)
        if len(found) > 1:
            anon_only = [f for f in found if f.stem.endswith("_anon")]
            if len(anon_only) == 1:
                found = anon_only
            else:
                print(f"Znaleziono wiele plikow: {[str(f) for f in found]}")
                print("Podaj recznie: python deanonymize.py input_anon.xlsx output.xlsx")
                sys.exit(1)

        input_file = found[0]
        stem = input_file.stem
        out_stem = stem[:-5] if stem.endswith("_anon") else stem + "_restored"
        output_file = out_stem + input_file.suffix
        print(f"Wykryto: {input_file} -> {output_file}")
        deanonymize(str(input_file), output_file)
