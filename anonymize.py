# -*- coding: utf-8 -*-
"""
anonymize.py - anonimizacja .xlsx na podstawie entities.xlsx.

Czyta entities.xlsx (analyze.py + reczna edycja).
Podmienia encje na pseudonimy (kolumna Pseudonim) tam gdzie Anonimizuj=TAK.
Bez LLM - czysty find-replace.

Uzycie:
    python anonymize.py                    (autodetekcja)
    python anonymize.py moj_plik.xlsx

Wymaga: entities.xlsx
Generuje: *_anon.xlsx + anon_map.json
"""

import os
import re
import sys
import json
import time
import glob
from typing import List, Dict, Tuple
from openpyxl import load_workbook

ENTITIES_FILE = "entities.xlsx"
MAP_FILE      = "anon_map.json"
GROUPS        = ["OSOBA", "FIRMA", "MIEJSCE", "DANE"]


def load_entities(path: str) -> List[Tuple[str, str, str]]:
    """
    Wczytuje entities.xlsx.
    Zwraca [(kategoria, wartosc, pseudonim), ...] dla wierszy z Anonimizuj=TAK.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    header = {}
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value:
            header[str(cell.value).strip()] = cell.column

    col_kat    = header.get("Kategoria")
    col_val    = header.get("Wartosc")
    col_pseudo = header.get("Pseudonim")
    col_anon   = header.get("Anonimizuj")

    if not all([col_kat, col_val, col_pseudo, col_anon]):
        wb.close()
        raise ValueError(
            f"Brak wymaganych kolumn w {path}. "
            f"Oczekiwano: Kategoria, Wartosc, Pseudonim, Anonimizuj. "
            f"Znaleziono: {list(header.keys())}"
        )

    entries: List[Tuple[str, str, str]] = []
    seen = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        try:
            kat    = str(row[col_kat    - 1] or "").strip()
            val    = str(row[col_val    - 1] or "").strip()
            pseudo = str(row[col_pseudo - 1] or "").strip()
            anon   = str(row[col_anon   - 1] or "").strip().upper()
        except IndexError:
            continue

        if anon == "TAK" and val and pseudo and val not in seen:
            entries.append((kat, val, pseudo))
            seen.add(val)

    wb.close()
    return entries


def process_workbook(in_path: str, out_path: str,
                     entries: List[Tuple[str, str, str]]) -> None:
    # Sortuj od najdluzszych wartosci (unika czesciowych podmian)
    entries_sorted = sorted(entries, key=lambda x: len(x[1]), reverse=True)

    # Buduj liste podmian i mape
    replacements: List[Tuple[str, str]] = []       # (wartosc, pseudonim)
    token_to_original: Dict[str, List[str]] = {}   # {pseudonim: [wartosc]}

    print(f"  Encji do anonimizacji: {len(entries_sorted)}")
    cat_counts: Dict[str, int] = {}
    for kat, val, pseudo in entries_sorted:
        replacements.append((val, pseudo))
        token_to_original[pseudo] = [val]
        cat_counts[kat] = cat_counts.get(kat, 0) + 1

    for cat, n in sorted(cat_counts.items()):
        print(f"    {cat}: {n}")

    print(f"  Laduje: {in_path}")
    t0 = time.time()
    wb = load_workbook(in_path)
    print(f"  Zaladowano w {time.time()-t0:.1f}s")

    original_cells: List[dict] = []
    changed = 0

    for sn in wb.sheetnames:
        ws = wb[sn]
        sheet_changed = 0
        for row in ws.iter_rows():
            for c in row:
                if not c.value or not isinstance(c.value, str):
                    continue
                original = c.value
                anonymized = original
                for val, pseudo in replacements:
                    anonymized = anonymized.replace(val, pseudo)
                if anonymized != original:
                    original_cells.append({
                        "sheet": sn, "row": c.row, "col": c.column,
                        "original": original, "anonymized": anonymized,
                    })
                    c.value = anonymized
                    sheet_changed += 1
        if sheet_changed:
            print(f"    {sn}: {sheet_changed} komorek zmienionych")
        changed += sheet_changed

    anon_map = {
        "token_to_original": token_to_original,
        "original_cells": original_cells,
    }
    with open(MAP_FILE, "w", encoding="utf-8") as f:
        json.dump(anon_map, f, ensure_ascii=False, indent=2)

    wb.save(out_path)
    print(f"\n  Zapisano: {out_path}")
    print(f"  Mapa:    {MAP_FILE}")
    print(f"  Zmieniono: {changed} komorek ({time.time()-t0:.1f}s)")
    print(f"\n  Deanonimizacja: python deanonymize.py")


def find_inputs() -> List[str]:
    return [f for f in sorted(glob.glob("*.xlsx"))
            if not f.endswith("_anon.xlsx")
            and not f.startswith("~$")
            and f != ENTITIES_FILE]


def main() -> int:
    inputs = sys.argv[1:] if len(sys.argv) > 1 else find_inputs()
    if not inputs:
        print("Nie znaleziono .xlsx do anonimizacji.")
        return 1
    if not os.path.exists(ENTITIES_FILE):
        print(f"BLAD: Nie znaleziono {ENTITIES_FILE}")
        print(f"      Najpierw: python analyze.py")
        return 1

    entries = load_entities(ENTITIES_FILE)
    if not entries:
        print(f"BLAD: Brak encji z Anonimizuj=TAK w {ENTITIES_FILE}.")
        return 1

    print(f"Wczytano {len(entries)} encji z {ENTITIES_FILE}")

    for in_path in inputs:
        base, ext = os.path.splitext(in_path)
        out_path = f"{base}_anon{ext}"
        print(f"\n{'='*60}")
        print(f"  {in_path} -> {out_path}")
        print(f"{'='*60}")
        try:
            process_workbook(in_path, out_path, entries)
        except Exception as e:
            print(f"[BLAD] {in_path}: {type(e).__name__}: {e}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
